"""Merge two filled annotation workbooks: Krippendorff's alpha (interval), adjudication list, human gold labels.
Usage: python src/merge_annotations.py annotation/EnHi-LMP_annotation_A.xlsx annotation/EnHi-LMP_annotation_B.xlsx
Writes: annotation/agreement.json, annotation/adjudicate.csv, data/enhi_lmp_human.jsonl
Then re-run experiments on human labels:  LMP_DATA=enhi_lmp_human.jsonl python src/run_experiments.py exp1  (etc.)"""
import sys, os, json, csv, itertools
import numpy as np
from openpyxl import load_workbook
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

def read(path):
    ws = load_workbook(path, data_only=True)["Pairs"]
    return {r[0]: (int(r[5]) if r[5] not in (None, "") else None) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}

def kripp_alpha_interval(units):
    """units: list of lists of numeric ratings (missing removed). Krippendorff's alpha, interval metric."""
    units = [u for u in units if len(u) >= 2]
    n = sum(len(u) for u in units)
    Do = sum(sum((a - b) ** 2 for a, b in itertools.permutations(u, 2)) / (len(u) - 1) for u in units) / n
    allv = [v for u in units for v in u]
    De = sum((a - b) ** 2 for a, b in itertools.permutations(allv, 2)) / (n * (n - 1))
    return 1 - Do / De if De > 0 else float("nan")

if __name__ == "__main__":
    A, B = read(sys.argv[1]), read(sys.argv[2])
    key = dict(csv.reader(open(os.path.join(ROOT, "annotation", "blind_key.csv"), encoding="utf-8")))
    key.pop("code", None)
    rows = {json.loads(l)["id"]: json.loads(l) for l in open(os.path.join(ROOT, "data", "enhi_lmp.jsonl"), encoding="utf-8")}
    units, adj, out = [], [], []
    for code, pid in key.items():
        a, b = A.get(code), B.get(code); vals = [v for v in (a, b) if v is not None]
        units.append(vals)
        if a is not None and b is not None and abs(a - b) >= 3: adj.append((code, pid, a, b))
        r = dict(rows[pid]); r["silver_score"] = r["score"]
        if vals:
            r["score"] = float(np.mean(vals)); r["annotations"] = vals
            r["decision"] = "Preserved" if r["score"] >= 7.5 else ("Partially Preserved" if r["score"] >= 4.5 else "Not Preserved")
        out.append(r)
    alpha = kripp_alpha_interval(units)
    both = [(A[c], B[c]) for c in key if A.get(c) is not None and B.get(c) is not None]
    rep = dict(krippendorff_alpha_interval=alpha, n_pairs_both=len(both),
               exact_agreement=float(np.mean([a == b for a, b in both])) if both else None,
               within_1=float(np.mean([abs(a - b) <= 1 for a, b in both])) if both else None,
               n_to_adjudicate=len(adj))
    sil = [r["silver_score"] for r in out if "annotations" in r]; hum = [r["score"] for r in out if "annotations" in r]
    if len(hum) > 2: rep["pearson_silver_vs_human"] = float(np.corrcoef(sil, hum)[0, 1])
    json.dump(rep, open(os.path.join(ROOT, "annotation", "agreement.json"), "w"), indent=1)
    with open(os.path.join(ROOT, "annotation", "adjudicate.csv"), "w", newline="", encoding="utf-8") as fh:
        csv.writer(fh).writerows([("code", "pair_id", "annotator_A", "annotator_B")] + adj)
    with open(os.path.join(ROOT, "data", "enhi_lmp_human.jsonl"), "w", encoding="utf-8") as fh:
        for r in out: fh.write(json.dumps(r, ensure_ascii=False) + "\n")
    print(json.dumps(rep, indent=1))
