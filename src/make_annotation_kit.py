"""Create the human-annotation workbook (one per annotator) for the 320 EnHi-LMP pairs.
Pairs are shuffled and given blind codes; gold/silver scores and error types are NOT shown to annotators."""
import json, os, random, csv
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
rows = [json.loads(l) for l in open(os.path.join(ROOT, "data", "enhi_lmp.jsonl"), encoding="utf-8")]
rng = random.Random(2026); order = rows[:]; rng.shuffle(order)
key = [(f"P{i+1:03d}", r["id"]) for i, r in enumerate(order)]
os.makedirs(os.path.join(ROOT, "annotation"), exist_ok=True)
with open(os.path.join(ROOT, "annotation", "blind_key.csv"), "w", newline="", encoding="utf-8") as fh:
    csv.writer(fh).writerows([("code", "pair_id")] + key)

RUBRIC = [
    ("10", "Identical legal effect. Only wording changed; every party, duty, permission, condition, exception, number and date is kept."),
    ("9", "Legal effect preserved; a purely stylistic nuance is lost."),
    ("8", "Preserved; a trivial, legally irrelevant detail is lost or slightly blurred."),
    ("7", "Partially preserved: a minor secondary detail is missing (e.g. ‘in writing’), core rule intact."),
    ("6", "Partially preserved: a secondary procedural detail or one item of a list is missing."),
    ("5", "Partially preserved: a secondary condition is dropped or scope is slightly widened/narrowed."),
    ("4", "Not preserved: an important condition is lost or scope materially changed."),
    ("3", "Not preserved: a core condition, exception, number, date or the strength of a duty is changed."),
    ("2", "Not preserved: the rule now gives a materially different outcome (wrong amount, deadline, party, or new duty)."),
    ("1", "Contradiction (rule reversed) or the simplification is unrelated to the original."),
]
hdr_fill = PatternFill("solid", fgColor="1B2A41"); thin = Side(style="thin", color="C9D3DE")

def build(annotator):
    wb = Workbook()
    ins = wb.active; ins.title = "Instructions"
    ins["A1"] = f"EnHi-LMP human annotation — Annotator {annotator}"; ins["A1"].font = Font(bold=True, size=14)
    lines = [
        "Task: for each pair, judge whether the SIMPLIFICATION preserves the LEGAL meaning of the ORIGINAL.",
        "Enter one integer from 1 to 10 in column F (LMP score). Use the rubric below. Do not discuss pairs with the other annotator.",
        "Judge legal effect only (who must/may do what, under which conditions, exceptions, amounts, dates). Ignore readability and style.",
        "Optional: column G — short reason when you give 7 or less (e.g. ‘deadline changed’, ‘exception removed’).",
        "Hindi pairs are judged with the same rubric. If you cannot read a pair, leave F empty and write ‘skip’ in G.",
    ]
    for i, l in enumerate(lines): ins.cell(row=3 + i, column=1, value=l)
    ins.cell(row=9, column=1, value="Score").font = Font(bold=True); ins.cell(row=9, column=2, value="Rubric anchor").font = Font(bold=True)
    for i, (s, d) in enumerate(RUBRIC):
        ins.cell(row=10 + i, column=1, value=int(s)); ins.cell(row=10 + i, column=2, value=d)
    ins.column_dimensions["A"].width = 10; ins.column_dimensions["B"].width = 120
    ws = wb.create_sheet("Pairs")
    heads = ["Code", "Language", "Domain", "Original", "Simplification", "LMP score (1–10)", "Reason (optional)"]
    for j, h in enumerate(heads, 1):
        c = ws.cell(row=1, column=j, value=h); c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, ((code, _), r) in enumerate(zip(key, order), 2):
        vals = [code, "English" if r["lang"] == "en" else "Hindi", r["domain"].title(), r["original"], r["simplification"], None, None]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v); c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if j == 6: c.fill = PatternFill("solid", fgColor="FBF3E6")
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="10", allow_blank=True,
                        showErrorMessage=True, errorTitle="Invalid score", error="Enter a whole number from 1 to 10.")
    ws.add_data_validation(dv); dv.add(f"F2:F{len(order)+1}")
    for col, w in zip("ABCDEFG", [8, 10, 11, 60, 60, 12, 30]): ws.column_dimensions[col].width = w
    ws.freeze_panes = "D2"; ws.auto_filter.ref = f"A1:G{len(order)+1}"
    prog = wb.create_sheet("Progress")
    prog["A1"] = "Scored pairs"; prog["B1"] = f"=COUNT(Pairs!F2:F{len(order)+1})"
    prog["A2"] = "Total pairs"; prog["B2"] = len(order)
    prog["A3"] = "Completion"; prog["B3"] = "=B1/B2"; prog["B3"].number_format = "0%"
    prog["A4"] = "Mean score"; prog["B4"] = f'=IFERROR(AVERAGE(Pairs!F2:F{len(order)+1}),"–")'
    for c in ("A1", "A2", "A3", "A4"): prog[c].font = Font(bold=True)
    prog.column_dimensions["A"].width = 16
    path = os.path.join(ROOT, "annotation", f"EnHi-LMP_annotation_{annotator}.xlsx"); wb.save(path); return path

for a in ["A", "B"]: print(build(a))
